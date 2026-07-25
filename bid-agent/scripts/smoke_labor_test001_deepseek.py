"""Per-node functional E2E test on a REAL .docx tender — REAL DeepSeek LLM.

Tender : 附件一：招标文件--2026年度SHEIN质检业务外包项目.docx (劳务外包/质检外包)
Bidder : 广州华南人力
Bid no : test001
Type   : 劳务外包类

This run uses the REAL DeepSeek API key from ~/.bid-agent/config.json so that
every LLM node (ReqExtractor / ContractExtractor / SectionGenerator /
QualityChecker / FeedbackProcessor / ScoreSimulator) produces REAL content.
Embeddings fall back to MockEmbedder (the configured OpenAI key is invalid),
so TemplateMatcher semantic retrieval is degraded but keyword fallback works.
"""
import sys
import os
import time
import json
import traceback
from pathlib import Path

sys.path.insert(0, "/Users/alanchris/Desktop/标书agent02/bid-agent")

from core.graph import run_pipeline, set_pipeline_llm, _LLM_NODES
from core.state import factory_state
from core.config_persistence import load_config
from core.llm.router import get_router
from core.llm.providers import test_connection
from langchain_core.messages import HumanMessage

DOC = "/Users/alanchris/Desktop/标书agent02/附件一：招标文件--2026年度SHEIN质检业务外包项目.docx"
BIDDER = "广州华南人力"
BID_NO = "test001"


def _flag(cond: bool) -> str:
    return "  <-- 异常" if cond else ""


def build_real_llm_fn(provider: str, model: str):
    """Build a (prompt: str) -> str callable backed by the real router."""
    router = get_router()

    def _fn(prompt: str) -> str:
        last = None
        for attempt in range(3):
            try:
                llm = router.get_llm("__pipeline__", provider=provider, model=model)
                resp = llm.invoke([HumanMessage(content=prompt)])
                content = resp.content if isinstance(resp.content, str) else str(resp.content)
                return content
            except Exception as e:  # noqa: BLE001
                last = e
                time.sleep(3 * (attempt + 1))
        # give up after retries — surface the error so the node sees a failure
        raise RuntimeError(f"DeepSeek call failed after retries: {last}") from last

    return _fn


def main() -> None:
    print("=" * 78)
    print("PER-NODE E2E (REAL DeepSeek LLM) — 劳务外包标书 (docx)")
    print(f"  tender : {Path(DOC).name}")
    print(f"  投标人 : {BIDDER}   标号 : {BID_NO}")
    print("=" * 78)

    cfg = load_config()
    ds_key = cfg["api_keys"].get("deepseek", "")
    print(f"[cfg] provider={cfg.get('provider')} model={cfg.get('model')} "
          f"deepseek_key_len={len(ds_key)}")

    # quick connectivity check (fail fast)
    ok, msg = test_connection("deepseek", ds_key, "deepseek-chat", timeout=20)
    print(f"[conn] DeepSeek: {'OK' if ok else 'FAIL'} — {msg}")
    if not ok:
        print("!! DeepSeek unavailable — aborting real run.")
        return

    # activate the real key on the global router
    router = get_router()
    router.set_api_key("deepseek", ds_key)
    # NOTE: OpenAI key in config is invalid (3 chars / 403) -> leave embeddings
    # to MockEmbedder (do NOT set OPENAI_API_KEY, avoids 403 errors).

    real_fn = build_real_llm_fn("deepseek", "deepseek-chat")
    set_pipeline_llm(real_fn)  # used by detect_subtype etc.
    llm_fns = {n: real_fn for n in _LLM_NODES}

    if not Path(DOC).exists():
        print("!! MISSING TENDER DOC:", DOC)
        return

    state = factory_state(
        documents=[{"filename": Path(DOC).name, "type": "docx", "path": DOC}],
        bid_type="劳务外包类",
        max_rounds=1,
        min_section_chars=0,
        project_contract={
            "project_name": f"{BID_NO}劳务外包项目",
            "bidder_name": BIDDER,
            "project_code": BID_NO,
            "industry": "劳务外包",
        },
        extra_reqs=f"投标人：{BIDDER}；招标编号：{BID_NO}；项目类型：劳务外包",
        requirements={
            "bid_number": BID_NO,
            "project_name": f"{BID_NO}劳务外包项目",
            "scoring": [], "qualifications": [], "tech_specs": [], "format_rules": {},
        },
    )
    state["review_status"] = "approved"

    t0 = time.time()
    try:
        result = run_pipeline(state, llm_fns=llm_fns)
    except Exception as e:  # noqa: BLE001
        print("!! EXCEPTION:", repr(e))
        traceback.print_exc()
        return
    print(f"\n[done] pipeline finished in {time.time()-t0:.1f}s")

    ns = result.get("node_status", {})
    print("\n########## 1) NODE STATUS ##########")
    for k, v in ns.items():
        print(f"  {k:22s}: {v}{_flag(v=='failed')}")
    print(f"  >> 失败节点: {[k for k,v in ns.items() if v=='failed'] or '无'}")

    # ---- 2) DocumentParser ----
    print("\n########## 2) DocumentParser ##########")
    for d in result.get("documents", []):
        n_tables = len(d.get("tables", [])) or len(d.get("extracted_tables", []))
        print(f"  status={d.get('status')}  chunks={len(d.get('chunks', []))}  tables={n_tables}{_flag(d.get('status')!='parsed')}")
        print(f"  text_len={len(d.get('content', '') or d.get('text', ''))}")

    # ---- 3) ReqExtractor ----
    print("\n########## 3) ReqExtractor (REAL) ##########")
    req = result.get("requirements", {})
    print(f"  project_name={req.get('project_name')!r}")
    print(f"  bid_number  ={req.get('bid_number')!r}  (预置被覆盖属预期)")
    print(f"  tenderer    ={req.get('tenderer_name')!r}")
    print(f"  bid_subtype ={result.get('bid_subtype')!r}")
    print(f"  scoring_items   = {len(req.get('scoring', []))}")
    print(f"  qualifications  = {len(req.get('qualifications', []))}")
    print(f"  tech_specs      = {len(req.get('tech_specs', []))}")
    print(f"  format_rules    = {list(req.get('format_rules', {}).keys())}")
    if req.get("scoring"):
        for s in req["scoring"][:3]:
            print("    scoring:", str(s)[:80])
    if req.get("qualifications"):
        for q in req["qualifications"][:3]:
            print("    qual   :", str(q)[:80])

    # ---- 4) ContractExtractor ----
    print("\n########## 4) ContractExtractor (REAL) ##########")
    pc = result.get("project_contract", {})
    print(f"  bidder_name ={pc.get('bidder_name')!r}")
    print(f"  project_name={pc.get('project_name')!r}")
    print(f"  project_code={pc.get('project_code')!r}")
    print(f"  industry    ={pc.get('industry')!r}")
    print(f"  tenderer    ={pc.get('tenderer_name')!r}")
    try:
        from core.contract import ProjectContextContract
        ok, errs = ProjectContextContract.from_dict(pc).is_valid()
        print(f"  is_valid={ok}  errors={errs}")
    except Exception as e:
        print("  contract validity check error:", repr(e))

    # ---- 5) InfoVerificationGate ----
    print("\n########## 5) InfoVerificationGate ##########")
    print(f"  info_verification_status={result.get('info_verification_status')!r}")
    isum = result.get("info_summary", {})
    print(f"  info_summary keys={list(isum.keys())[:10]}")

    # ---- 6) EligibilityChecker ----
    print("\n########## 6) EligibilityChecker ##########")
    er = result.get("eligibility_report", {})
    print(f"  verdict={er.get('verdict')!r}")
    print(f"  missing_quals={len(er.get('missing_quals', []))}  present_quals={len(er.get('present_quals', []))}")

    # ---- 7) TemplateMatcher ----
    print("\n########## 7) TemplateMatcher ##########")
    print(f"  matched_templates  ={result.get('matched_templates')}")
    print(f"  selected_template  ={result.get('selected_template_id')!r}")

    # ---- 8) SectionGenerator ----
    print("\n########## 8) SectionGenerator (REAL) ##########")
    secs = result.get("sections", {})
    print(f"  section count = {len(secs)}{_flag(len(secs)==0)}")
    mock_like = 0
    for name, c in secs.items():
        clen = len(c) if isinstance(c, str) else -1
        if isinstance(c, str) and 70 <= clen <= 95:
            mock_like += 1
        sample = (c[:60] + "…") if isinstance(c, str) and c else str(c)[:60]
        print(f"    - {name}: {clen} chars | {sample}")
    print(f"  >> mock占位-like 章节数: {mock_like}  (0 表示已生成真实内容)")
    # total real content length
    total_chars = sum(len(c) for c in secs.values() if isinstance(c, str))
    print(f"  >> 全部章节总字数: {total_chars}")

    # ---- 9) QualityChecker ----
    print("\n########## 9) QualityChecker ##########")
    qr = result.get("quality_report", {})
    print(f"  verdict={qr.get('verdict')!r}  total={qr.get('total_items')}  passed={qr.get('passed_items')}")
    print(f"  completeness keys={list(qr.get('completeness', {}).keys())}")
    print(f"  compliance issues={len(qr.get('compliance', []))}  consistency issues={len(qr.get('consistency', []))}")
    if qr.get('compliance'):
        for x in qr['compliance'][:3]:
            print("    issue:", str(x)[:80])

    # ---- 10) FeedbackProcessor ----
    print("\n########## 10) FeedbackProcessor ##########")
    print(f"  current_round={result.get('current_round')}  status(pending正常)={ns.get('FeedbackProcessor')}")

    # ---- 11) CrossReferenceChecker ----
    print("\n########## 11) CrossReferenceChecker ##########")
    cr = result.get("cross_ref_report", {})
    print(f"  keys={list(cr.keys())}  issues={len(cr.get('issues', [])) if isinstance(cr, dict) else 'n/a'}")

    # ---- 12) ComplianceChecker ----
    print("\n########## 12) ComplianceChecker ##########")
    cm = result.get("compliance_report", {})
    print(f"  type={type(cm).__name__}  keys={list(cm.keys()) if isinstance(cm, dict) else 'n/a'}")
    if isinstance(cm, dict):
        print(f"  verdict={cm.get('verdict')!r}  issues={len(cm.get('issues', []))}")

    # ---- 13) ScoreSimulator ----
    print("\n########## 13) ScoreSimulator ##########")
    ss = result.get("score_simulation", {})
    print(f"  type={type(ss).__name__}  keys={list(ss.keys()) if isinstance(ss, dict) else 'n/a'}")
    if isinstance(ss, dict):
        print(f"  total={ss.get('total')}  items={len(ss.get('items', []))}")
        for it in (ss.get('items', []) or [])[:5]:
            print("    score item:", str(it)[:80])

    # ---- 14) HumanReviewGate ----
    print("\n########## 14) HumanReviewGate ##########")
    print(f"  review_status={result.get('review_status')!r}")

    # ---- 15) DocumentAssembler ----
    print("\n########## 15) DocumentAssembler ##########")
    ep = result.get("export_path")
    print(f"  export_path={ep}")
    if ep and Path(ep).exists():
        _inspect_docx(ep)
    else:
        print("  !! 未生成导出文件")

    print("\nE2E_DONE")


def _inspect_docx(path: str) -> None:
    try:
        from docx import Document
        doc = Document(path)
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        full = "\n".join(paras)
        print(f"  docx 非空段落={len(paras)} 表格={len(doc.tables)}")
        print(f"  contains 广州华南人力 : {'广州华南人力' in full}")
        print(f"  contains test001      : {'test001' in full}")
        # estimate real content: count chars
        print(f"  total text chars      : {len(full)}")
        print("  --- 前 12 非空段落 ---")
        for t in paras[:12]:
            print("   |", t[:90])
    except Exception as e:  # noqa: BLE001
        print("  !! docx 检查错误:", repr(e))


if __name__ == "__main__":
    main()
