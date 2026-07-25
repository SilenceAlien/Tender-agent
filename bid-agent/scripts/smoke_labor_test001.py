"""E2E full-pipeline test: simulate a labor-outsourcing bid submission.

Bidder : 广州华南人力
Bid no : test001
Type   : 劳务外包类

No API key in env -> LLM nodes default to mock.  This exercises the REAL
DocumentParser / ReqExtractor / SectionGenerator / DocumentAssembler path on a
real tender PDF (not the mock-PDF unit tests), and verifies the F3/F4 fixes
(no silent eligibility dead-end, no 30k truncation) keep the pipeline flowing
end-to-end.

Submission metadata is injected through the legitimate "human-confirmed"
channels (project_contract + requirements) so we can confirm bidder/bid-no
survive the pipeline and reach the export filename + document body.
"""
import sys
import traceback
from pathlib import Path

sys.path.insert(0, "/Users/alanchris/Desktop/标书agent02/bid-agent")

from core.graph import run_pipeline
from core.state import factory_state

# Real labor-outsourcing tender doc at project root
DOC = "/Users/alanchris/Desktop/标书agent02/2026年广州铁道车辆有限公司劳务管理服务采购（技术标）v2.pdf"

BIDDER = "广州华南人力"
BID_NO = "test001"


def main() -> None:
    print("=" * 72)
    print(f"E2E TEST — 劳务外包标书  投标人={BIDDER}  标号={BID_NO}")
    print("=" * 72)

    if not Path(DOC).exists():
        print("!! MISSING TENDER DOC:", DOC)
        return

    state = factory_state(
        documents=[{"filename": Path(DOC).name, "type": "pdf", "path": DOC}],
        bid_type="劳务外包类",
        max_rounds=1,
        min_section_chars=0,
        # Simulate human-confirmed key info (InfoVerificationGate path)
        project_contract={
            "project_name": f"{BID_NO}劳务外包项目",
            "bidder_name": BIDDER,
            "project_code": BID_NO,
            "industry": "劳务外包",
        },
        extra_reqs=f"投标人：{BIDDER}；招标编号：{BID_NO}；项目类型：劳务外包",
        requirements={
            "bid_number": BID_NO,
            "project_name": f"{BID_NO}劳务外包项目",
            "scoring": [],
            "qualifications": [],
            "tech_specs": [],
            "format_rules": {},
        },
    )
    # Bypass the human review gate for headless run
    state["review_status"] = "approved"

    try:
        result = run_pipeline(state)
    except Exception as e:  # noqa: BLE001
        print("!! EXCEPTION:", repr(e))
        traceback.print_exc()
        return

    ns = result.get("node_status", {})
    print("\n[NODE STATUS]")
    for k, v in ns.items():
        print(f"  {k:22s}: {v}")
    failed = [k for k, v in ns.items() if v == "failed"]
    print(f"\n  FAILED NODES : {failed or 'NONE'}")

    docs = result.get("documents", [])
    for d in docs:
        n_tables = len(d.get("tables", [])) or len(d.get("extracted_tables", []))
        print(f"\n[DOC {d.get('filename')}] status={d.get('status')} "
              f"chunks={len(d.get('chunks', []))} tables={n_tables}")

    secs = result.get("sections", {})
    print(f"\n[SECTIONS] count={len(secs)}")
    for name, c in list(secs.items()):
        print(f"  {name}: {len(c) if isinstance(c, str) else type(c).__name__} chars")

    pc = result.get("project_contract", {})
    req = result.get("requirements", {})
    print(f"\n[CONTRACT] bidder_name={pc.get('bidder_name')!r} "
          f"project_name={pc.get('project_name')!r}")
    print(f"[REQ]     bid_number={req.get('bid_number')!r}")

    ep = result.get("export_path")
    print(f"\n[EXPORT] {ep}")
    if ep and Path(ep).exists():
        _inspect_docx(ep)
    else:
        print("  !! no export file produced")


def _inspect_docx(path: str) -> None:
    try:
        from docx import Document
        doc = Document(path)
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        full = "\n".join(paras)
        print(f"  docx non-empty paragraphs={len(paras)} tables={len(doc.tables)}")
        print(f"  contains 广州华南人力 : {'广州华南人力' in full}")
        print(f"  contains test001      : {'test001' in full}")
        print("  --- first 10 non-empty paragraphs ---")
        for t in paras[:10]:
            print("   |", t[:90])
    except Exception as e:  # noqa: BLE001
        print("  !! docx inspect error:", repr(e))


if __name__ == "__main__":
    main()
    print("\nE2E_DONE")
