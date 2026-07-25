"""Integration tests for Markdown formatting support in DOCX export.

Verifies:
- Markdown tables are rendered as Word tables
- Markdown headings (##, ###) are rendered with appropriate fonts/sizes
- Markdown list items (-, *, 1.) are rendered with bullets/numbers
- **bold** text is rendered as bold runs
- <br> tags are converted to line breaks
- Mermaid edge label syntax (-- label -->) is parsed correctly
- Mermaid & multi-node syntax is parsed correctly
- Mermaid labels with HTML tags are cleaned correctly
"""

import tempfile
from pathlib import Path

import pytest

from core.nodes.doc_assembler import (
    _build_docx,
    _mermaid_to_text_diagram,
    _clean_mermaid_label,
    _is_markdown_table_line,
    _is_separator_line,
    _parse_markdown_table_row,
    DEFAULT_FORMAT,
)


# ── Markdown Table Helper Tests ────────────────────────────────────────


class TestMarkdownTableHelpers:
    """Tests for markdown table parsing helpers."""

    def test_is_markdown_table_line_valid(self):
        assert _is_markdown_table_line("| A | B | C |")
        assert _is_markdown_table_line("|名称|数量|")
        assert _is_markdown_table_line("| 1 | 2 |")

    def test_is_markdown_table_line_invalid(self):
        assert not _is_markdown_table_line("普通文本")
        assert not _is_markdown_table_line("|单列")
        assert not _is_markdown_table_line("单列|")
        assert not _is_markdown_table_line("")

    def test_is_separator_line_valid(self):
        assert _is_separator_line("|---|---|---|")
        assert _is_separator_line("|:--|--:|:--:|")
        assert _is_separator_line("| --- | --- |")

    def test_is_separator_line_invalid(self):
        assert not _is_separator_line("| A | B |")
        assert not _is_separator_line("|--|abc|")
        assert not _is_separator_line("普通文本")

    def test_parse_table_row(self):
        cells = _parse_markdown_table_row("| 名称 | 数量 | 单价 |")
        assert cells == ["名称", "数量", "单价"]

    def test_parse_table_row_with_escaped_pipe(self):
        cells = _parse_markdown_table_row("| 说明 | A\\|B |")
        assert cells == ["说明", "A|B"]

    def test_parse_table_row_single_cell(self):
        cells = _parse_markdown_table_row("| 内容 |")
        assert cells == ["内容"]


# ── DOCX Table Export Tests ────────────────────────────────────────────


class TestDocxTableExport:
    """Tests for Markdown table → Word table conversion."""

    def test_markdown_table_becomes_word_table(self):
        """Markdown table should be rendered as a Word table, not plain text."""
        sections = {
            "技术方案": (
                "设备清单如下：\n"
                "| 名称 | 数量 | 规格 |\n"
                "|------|------|------|\n"
                "| 服务器 | 5 | 2U |\n"
                "| 交换机 | 3 | 24口 |\n"
                "以上为设备清单。"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_table.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)

            # Should have at least one table
            assert len(doc.tables) >= 1, "Expected at least one Word table"

            table = doc.tables[0]
            # Header row
            assert table.rows[0].cells[0].text.strip() == "名称"
            assert table.rows[0].cells[1].text.strip() == "数量"
            assert table.rows[0].cells[2].text.strip() == "规格"

            # Data rows
            assert table.rows[1].cells[0].text.strip() == "服务器"
            assert table.rows[1].cells[1].text.strip() == "5"
            assert table.rows[1].cells[2].text.strip() == "2U"
            assert table.rows[2].cells[0].text.strip() == "交换机"
            assert table.rows[2].cells[1].text.strip() == "3"
            assert table.rows[2].cells[2].text.strip() == "24口"

    def test_table_with_placeholder_in_cell(self):
        """Placeholders inside table cells should be rendered in red bold."""
        sections = {
            "报价方案": (
                "| 项目 | 金额 |\n"
                "|------|------|\n"
                "| 设备费 | 【待填写：设备金额】 |\n"
                "| 服务费 | 【待填写：服务金额】 |\n"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_table_placeholder.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            from docx.shared import RGBColor
            doc = Document(export_path)

            assert len(doc.tables) >= 1
            table = doc.tables[0]

            # Find placeholder runs in table cells
            placeholder_runs = []
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            if "待填写" in run.text:
                                placeholder_runs.append(run)

            assert len(placeholder_runs) >= 2
            for run in placeholder_runs:
                assert run.bold is True
                assert run.font.color.rgb == RGBColor(0xFF, 0x00, 0x00)

    def test_multiple_tables_in_section(self):
        """Multiple markdown tables in one section should all be converted."""
        sections = {
            "技术方案": (
                "表1：\n"
                "| A | B |\n"
                "|---|---|\n"
                "| 1 | 2 |\n"
                "\n"
                "表2：\n"
                "| C | D |\n"
                "|---|---|\n"
                "| 3 | 4 |\n"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_multi_table.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)
            assert len(doc.tables) >= 2

    def test_table_text_not_in_paragraphs(self):
        """Table content should not also appear as raw paragraph text."""
        sections = {
            "测试": (
                "| 名称 | 数量 |\n"
                "|------|------|\n"
                "| 电脑 | 10 |\n"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_no_raw.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)
            all_text = " ".join(p.text for p in doc.paragraphs)

            # The pipe characters should not appear in paragraph text
            assert "|---|" not in all_text
            assert "| 名称 |" not in all_text


# ── Markdown Heading Tests ─────────────────────────────────────────────


class TestDocxMarkdownHeading:
    """Tests for Markdown heading export."""

    def test_heading_exported_with_bold(self):
        """## heading should be rendered as bold text."""
        sections = {
            "技术方案": (
                "## 1. 系统架构\n"
                "系统采用三层架构。\n"
                "### 1.1 前端层\n"
                "使用Vue.js框架。"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_heading.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)
            all_text = "\n".join(p.text for p in doc.paragraphs)

            assert "系统架构" in all_text
            assert "前端层" in all_text

            # Find heading paragraphs — they should have bold runs
            heading_found = False
            for para in doc.paragraphs:
                if "系统架构" in para.text:
                    for run in para.runs:
                        if run.text.strip() and run.bold:
                            heading_found = True
                            break
            assert heading_found, "Heading text should be bold"

    def test_heading_no_hash_in_output(self):
        """## markers should not appear in the output text."""
        sections = {
            "测试": "## 标题\n正文内容。",
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_no_hash.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)
            all_text = "\n".join(p.text for p in doc.paragraphs)

            assert "##" not in all_text
            assert "标题" in all_text


# ── Markdown List Tests ────────────────────────────────────────────────


class TestDocxMarkdownList:
    """Tests for Markdown list export."""

    def test_unordered_list_exported(self):
        """List items with - or * should be rendered with bullet points."""
        sections = {
            "服务方案": (
                "服务内容：\n"
                "- 日常维护\n"
                "- 故障排除\n"
                "- 定期巡检\n"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_list.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)
            all_text = "\n".join(p.text for p in doc.paragraphs)

            assert "日常维护" in all_text
            assert "故障排除" in all_text
            assert "定期巡检" in all_text
            # Should have bullet characters
            assert "•" in all_text

    def test_ordered_list_exported(self):
        """Numbered list items should preserve their numbers."""
        sections = {
            "实施计划": (
                "实施步骤：\n"
                "1. 需求分析\n"
                "2. 方案设计\n"
                "3. 系统部署\n"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_ordered_list.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)
            all_text = "\n".join(p.text for p in doc.paragraphs)

            assert "需求分析" in all_text
            assert "方案设计" in all_text
            assert "系统部署" in all_text

    def test_list_no_dash_marker_in_output(self):
        """- markers should not appear as raw text (converted to •)."""
        sections = {
            "测试": "- 项目一\n- 项目二",
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_no_dash.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)

            # Check that no paragraph starts with "- " (raw dash)
            for para in doc.paragraphs:
                assert not para.text.strip().startswith("- "), (
                    f"Raw dash marker found: '{para.text}'"
                )


# ── Bold Text Tests ────────────────────────────────────────────────────


class TestDocxBoldText:
    """Tests for **bold** markdown export."""

    def test_bold_text_rendered(self):
        """**bold** text should be rendered as bold runs."""
        sections = {
            "技术方案": "本系统采用**微服务架构**，确保高可用性。",
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_bold.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)

            # Find the paragraph with bold content
            bold_runs = []
            for para in doc.paragraphs:
                if "微服务架构" in para.text:
                    for run in para.runs:
                        if "微服务架构" in run.text:
                            bold_runs.append(run)

            assert len(bold_runs) >= 1
            for run in bold_runs:
                assert run.bold is True, (
                    f"'{run.text}' should be bold"
                )

    def test_bold_markers_not_in_output(self):
        """** markers should not appear in the output text."""
        sections = {
            "测试": "这是**重要**内容。",
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_no_stars.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)
            all_text = " ".join(p.text for p in doc.paragraphs)

            assert "**" not in all_text
            assert "重要" in all_text

    def test_bold_and_placeholder_coexist(self):
        """**bold** and 【待填写：...】 should coexist in the same paragraph."""
        sections = {
            "测试": "项目编号**非常重要**：【待填写：项目编号】请确认。",
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_bold_placeholder.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            from docx.shared import RGBColor
            doc = Document(export_path)

            # Find the paragraph
            target_para = None
            for para in doc.paragraphs:
                if "非常重要" in para.text and "待填写" in para.text:
                    target_para = para
                    break

            assert target_para is not None

            # Should have bold runs for "非常重要"
            bold_runs = [r for r in target_para.runs if "非常重要" in r.text]
            assert len(bold_runs) >= 1
            assert all(r.bold for r in bold_runs)

            # Should have red bold runs for placeholder
            placeholder_runs = [r for r in target_para.runs if "待填写" in r.text]
            assert len(placeholder_runs) >= 1
            for r in placeholder_runs:
                assert r.bold is True
                assert r.font.color.rgb == RGBColor(0xFF, 0x00, 0x00)


# ── <br> Tag Tests ─────────────────────────────────────────────────────


class TestDocxBreakTag:
    """Tests for <br> tag handling."""

    def test_br_converted_to_line_break(self):
        """<br> tags should be converted to line breaks within a paragraph."""
        sections = {
            "测试": "第一行<br>第二行<br>第三行",
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_br.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)

            # Find the paragraph with the content
            target_para = None
            for para in doc.paragraphs:
                if "第一行" in para.text:
                    target_para = para
                    break

            assert target_para is not None
            # <br> should not appear in text
            assert "<br>" not in target_para.text
            assert "<br" not in target_para.text.lower()
            # All three lines should be present
            assert "第一行" in target_para.text
            assert "第二行" in target_para.text
            assert "第三行" in target_para.text


# ── Mermaid Enhancement Tests ──────────────────────────────────────────


class TestMermaidEnhancements:
    """Tests for enhanced Mermaid parsing features."""

    def test_clean_mermaid_label_strips_html(self):
        """HTML tags in labels should be stripped."""
        assert _clean_mermaid_label("hello<br>world") == "hello world"
        assert _clean_mermaid_label("<b>bold</b>") == "bold"
        assert _clean_mermaid_label("no tags") == "no tags"

    def test_dash_label_syntax(self):
        """Edge labels using -- label --> syntax should be parsed."""
        mermaid_code = """
flowchart LR
    A[开始] -- 审核 --> B[执行]
"""
        lines = _mermaid_to_text_diagram(mermaid_code)
        joined = "\n".join(lines)
        assert "开始" in joined
        assert "执行" in joined
        assert "审核" in joined

    def test_amp_multi_node_syntax(self):
        """& multi-node edges should be parsed into individual edges."""
        mermaid_code = """
flowchart TD
    A[根节点] & B[节点B] --> C[目标]
"""
        lines = _mermaid_to_text_diagram(mermaid_code)
        joined = "\n".join(lines)
        assert "根节点" in joined
        assert "节点B" in joined
        assert "目标" in joined

    def test_html_tags_in_node_labels(self):
        """HTML tags like <br> in node labels should be cleaned."""
        mermaid_code = """
flowchart TD
    A[第一行<br>第二行] --> B[目标]
"""
        lines = _mermaid_to_text_diagram(mermaid_code)
        joined = "\n".join(lines)
        # <br> should not appear in the output
        assert "<br>" not in joined
        assert "第一行" in joined
        assert "第二行" in joined
        assert "目标" in joined

    def test_subgraph_ignored(self):
        """subgraph declarations should be skipped without crashing."""
        mermaid_code = """
flowchart TD
    subgraph Group1
        A[节点A] --> B[节点B]
    end
    B --> C[节点C]
"""
        lines = _mermaid_to_text_diagram(mermaid_code)
        joined = "\n".join(lines)
        assert "节点A" in joined
        assert "节点B" in joined
        assert "节点C" in joined


# ── Combined Format Tests ──────────────────────────────────────────────


class TestCombinedFormats:
    """Tests for combined Markdown features in a single section."""

    def test_table_heading_list_combined(self):
        """Tables, headings, and lists should all work in one section."""
        sections = {
            "技术方案": (
                "## 1. 系统概述\n"
                "本系统包含以下组件：\n"
                "- 应用服务器\n"
                "- 数据库服务器\n"
                "\n"
                "## 2. 设备清单\n"
                "| 设备 | 数量 |\n"
                "|------|------|\n"
                "| 服务器 | **10** |\n"
                "| 交换机 | 5 |\n"
                "\n"
                "```mermaid\n"
                "flowchart TD\n"
                "    A[核心交换机] --> B[服务器集群]\n"
                "```\n"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_combined.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)
            all_text = "\n".join(p.text for p in doc.paragraphs)

            # Headings
            assert "系统概述" in all_text
            assert "设备清单" in all_text
            assert "##" not in all_text

            # List items
            assert "应用服务器" in all_text
            assert "数据库服务器" in all_text
            assert "•" in all_text

            # Table
            assert len(doc.tables) >= 1
            table = doc.tables[0]
            assert table.rows[0].cells[0].text.strip() == "设备"
            assert table.rows[0].cells[1].text.strip() == "数量"
            assert table.rows[1].cells[0].text.strip() == "服务器"

            # Mermaid diagram — rendered as embedded image, not text
            assert "```mermaid" not in all_text
            assert len(doc.inline_shapes) >= 1, (
                f"Expected ≥1 inline image for mermaid diagram, got {len(doc.inline_shapes)}"
            )


# ── TOC Format Tests ───────────────────────────────────────────────────


class TestDocxTOCFormat:
    """Tests for Table of Contents formatting with tab stops."""

    def test_toc_uses_tab_not_dots(self):
        """TOC entries should use tab character, not hardcoded dot fills.

        F2 fix: H7 changed TOC page numbers from literal "页" text to PAGEREF
        fields.  Old test searched for paragraphs containing "页", which no
        longer matches.  Now we find TOC entries by section name + tab char.
        """
        sections = {
            "第一章 投标函": "内容",
            "第二章 技术方案": "内容",
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_toc.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)

            # F2 fix: find TOC entries by section name + tab character,
            # not by literal "页" (PAGEREF fields don't contain "页").
            toc_entries = [
                p for p in doc.paragraphs
                if "第" in p.text and "\t" in p.text
            ]
            assert len(toc_entries) >= 2

            for entry in toc_entries:
                # Should contain a tab character (\t) separating title from page
                assert "\t" in entry.text, (
                    f"TOC entry should contain tab: '{entry.text}'"
                )
                # Should NOT contain long strings of dots
                assert "··" not in entry.text
                assert "·" * 5 not in entry.text

    def test_toc_has_tab_stops(self):
        """TOC paragraphs should have right-aligned tab stops with dotted leaders.

        F2 fix: find TOC entry by section name + tab character instead of
        literal "页" (PAGEREF fields don't contain "页" text).
        """
        sections = {
            "第一章 测试": "内容",
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_toc_stops.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)

            # F2 fix: find TOC entry by section name + tab character
            toc_entry = None
            for p in doc.paragraphs:
                if "第一章" in p.text and "\t" in p.text:
                    toc_entry = p
                    break

            assert toc_entry is not None
            # Should have at least one tab stop
            assert len(toc_entry.paragraph_format.tab_stops) >= 1


# ── Diagram Image Rendering Tests ──────────────────────────────────────


class TestDiagramImageRendering:
    """Tests for Mermaid diagram → PNG image rendering in DOCX."""

    def test_simple_diagram_produces_image(self):
        """A simple flowchart should produce an embedded image in the DOCX."""
        sections = {
            "组织架构": (
                "```mermaid\n"
                "flowchart TD\n"
                "    A[总经理] --> B[技术部]\n"
                "    A --> C[市场部]\n"
                "    A --> D[财务部]\n"
                "```"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_diagram_img.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)

            assert len(doc.inline_shapes) >= 1
            # Image should have non-zero width and height
            shape = doc.inline_shapes[0]
            assert shape.width > 0
            assert shape.height > 0

    def test_lr_diagram_produces_image(self):
        """A left-to-right flowchart should produce an embedded image."""
        sections = {
            "流程": (
                "```mermaid\n"
                "flowchart LR\n"
                "    A[开始] --> B[处理]\n"
                "    B --> C[结束]\n"
                "```"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_lr_diagram.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)
            assert len(doc.inline_shapes) >= 1

    def test_diagram_with_edge_labels(self):
        """A flowchart with edge labels should produce an image."""
        sections = {
            "审批流程": (
                "```mermaid\n"
                "flowchart TD\n"
                "    A[提交] -->|审核通过| B[执行]\n"
                "    A -->|审核拒绝| C[退回]\n"
                "```"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_edge_label_img.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)
            assert len(doc.inline_shapes) >= 1

    def test_diagram_image_centered(self):
        """Diagram images should be in centered paragraphs."""
        sections = {
            "架构": (
                "```mermaid\n"
                "flowchart TD\n"
                "    A[根] --> B[子]\n"
                "```"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_centered.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            doc = Document(export_path)

            # Find the paragraph containing the image
            image_para = None
            for p in doc.paragraphs:
                if len(p.runs) > 0:
                    for run in p.runs:
                        if run._element.findall(
                            ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"
                        ):
                            image_para = p
                            break
                    if image_para:
                        break

            if image_para:
                assert image_para.alignment == WD_ALIGN_PARAGRAPH.CENTER

    def test_diagram_with_multiline_labels(self):
        """A flowchart with <br> in labels should still render as image."""
        sections = {
            "架构": (
                "```mermaid\n"
                "flowchart TD\n"
                "    A[第一行<br>第二行] --> B[目标]\n"
                "```"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_multiline_img.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)
            assert len(doc.inline_shapes) >= 1

    def test_text_fallback_on_invalid_mermaid(self):
        """If mermaid code has no edges, should fall back to text diagram."""
        sections = {
            "测试": (
                "```mermaid\n"
                "flowchart TD\n"
                "    A[孤立节点]\n"
                "```"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_fallback.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            doc = Document(export_path)
            all_text = "\n".join(p.text for p in doc.paragraphs)

            # Should have fallback text (no image since no edges)
            assert "```mermaid" not in all_text
