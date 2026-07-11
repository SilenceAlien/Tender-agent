"""Integration tests for DocumentAssembler node (T019 + T020)."""

import tempfile
from pathlib import Path

import pytest

from core.nodes.doc_assembler import (
    _build_docx,
    _cm_to_emu,
    _format_margin,
    DEFAULT_FORMAT,
    doc_assembler,
)
from core.state import NodeStatus, factory_state


# ── Fixtures ───────────────────────────────────────────────────────────


@pytest.fixture
def state_with_sections():
    """Full state with 8 sections and format rules."""
    return factory_state(
        requirements={
            "scoring": [{"item_name": "方案", "score": 50}],
            "format_rules": {
                "page_margin": "上3.7cm/下3.5cm/左2.8cm/右2.6cm",
                "font": "正文仿宋_GB2312",
                "line_spacing": "28磅",
                "title_levels": ["黑体 二号", "楷体 三号", "仿宋 四号"],
                "seal_requirement": "逐页盖章",
                "binding": "A4胶装",
            },
        },
        sections={
            "第一章 投标函": "投标函内容。\n我们郑重承诺按照招标文件要求提供全部服务。" * 5,
            "第二章 法定代表人授权委托书": "授权委托书内容。\n兹授权XXX代表本公司。" * 5,
            "第三章 服务方案": "服务方案详细内容。\n覆盖全部评分项，详尽论述。" * 5,
            "第四章 技术方案": "技术方案详细内容。\n完全满足技术规格要求。" * 5,
            "第五章 人员配置方案": "人员配置详细内容。\n团队10人，含PMP项目经理。" * 5,
            "第六章 公司资质与业绩": "资质业绩详细内容。\nISO9001认证，3年同类项目经验。" * 5,
            "第七章 项目实施计划": "实施计划详细内容。\n6个月项目周期，分4个阶段。" * 5,
            "第八章 售后服务承诺": "售后服务详细内容。\n7x24小时，30分钟响应。" * 5,
        },
        current_round=0,
    )


# ── Unit Tests ─────────────────────────────────────────────────────────


class TestFormatMargin:
    def test_standard_format(self):
        margins = _format_margin({"page_margin": "上3.7cm/下3.5cm/左2.8cm/右2.6cm"})
        assert margins["top"] == 3.7
        assert margins["bottom"] == 3.5
        assert margins["left"] == 2.8
        assert margins["right"] == 2.6

    def test_partial_format(self):
        margins = _format_margin({"page_margin": "上2.5cm/下2.0cm"})
        assert margins["top"] == 2.5
        assert margins["bottom"] == 2.0
        # Left/right use defaults
        assert margins["left"] == 2.8
        assert margins["right"] == 2.6

    def test_default_format(self):
        margins = _format_margin({})
        assert margins["top"] == 3.7
        assert margins["bottom"] == 3.5
        assert margins["left"] == 2.8
        assert margins["right"] == 2.6


class TestCmToEmu:
    def test_conversion(self):
        assert _cm_to_emu(1.0) == 360000
        assert _cm_to_emu(3.7) == pytest.approx(1332000)
        assert _cm_to_emu(2.8) == pytest.approx(1008000)


# ── Integration Tests ──────────────────────────────────────────────────


class TestBuildDocx:
    def test_creates_docx_file(self, state_with_sections):
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_output.docx")
            result = _build_docx(
                state_with_sections["sections"],
                state_with_sections["requirements"]["format_rules"],
                export_path,
            )
            assert result == export_path
            assert Path(export_path).exists()
            assert Path(export_path).stat().st_size > 1000  # At least 1KB

    def test_docx_has_all_sections(self, state_with_sections):
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_output.docx")

            from docx import Document
            _build_docx(
                state_with_sections["sections"],
                state_with_sections["requirements"]["format_rules"],
                export_path,
            )

            doc = Document(export_path)
            # Extract all paragraphs text
            all_text = " ".join(p.text for p in doc.paragraphs)

            assert "第一章 投标函" in all_text
            assert "第八章 售后服务承诺" in all_text
            assert "投标函内容" in all_text

    def test_docx_has_title(self, state_with_sections):
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_output.docx")
            _build_docx(
                state_with_sections["sections"],
                state_with_sections["requirements"]["format_rules"],
                export_path,
            )

            from docx import Document
            doc = Document(export_path)
            titles = [p.text for p in doc.paragraphs if p.text.strip()]
            assert any("投  标  书" in t for t in titles)

    def test_docx_has_toc(self, state_with_sections):
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_output.docx")
            _build_docx(
                state_with_sections["sections"],
                state_with_sections["requirements"]["format_rules"],
                export_path,
            )

            from docx import Document
            doc = Document(export_path)
            all_text = " ".join(p.text for p in doc.paragraphs)
            assert "目  录" in all_text

    def test_margins_applied(self, state_with_sections):
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_output.docx")
            _build_docx(
                state_with_sections["sections"],
                state_with_sections["requirements"]["format_rules"],
                export_path,
            )

            from docx import Document
            from docx.shared import Cm
            doc = Document(export_path)

            for section in doc.sections:
                # Cm uses integer EMU, so compare via approximate EMU values
                assert abs(section.top_margin - Cm(3.7)) < 1000
                assert abs(section.bottom_margin - Cm(3.5)) < 1000
                assert abs(section.left_margin - Cm(2.8)) < 1000
                assert abs(section.right_margin - Cm(2.6)) < 1000

    def test_line_spacing_applied(self, state_with_sections):
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_output.docx")
            _build_docx(
                state_with_sections["sections"],
                state_with_sections["requirements"]["format_rules"],
                export_path,
            )

            from docx import Document
            from docx.shared import Pt
            doc = Document(export_path)
            style = doc.styles["Normal"]
            assert style.paragraph_format.line_spacing == Pt(28)

    def test_first_line_indent_applied(self, state_with_sections):
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_output.docx")
            _build_docx(
                state_with_sections["sections"],
                state_with_sections["requirements"]["format_rules"],
                export_path,
            )

            from docx import Document
            doc = Document(export_path)
            style = doc.styles["Normal"]
            assert style.paragraph_format.first_line_indent is not None

    def test_seal_text_included(self, state_with_sections):
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_output.docx")
            _build_docx(
                state_with_sections["sections"],
                state_with_sections["requirements"]["format_rules"],
                export_path,
            )

            from docx import Document
            doc = Document(export_path)
            all_text = " ".join(p.text for p in doc.paragraphs)
            assert "逐页盖章" in all_text

    def test_placeholder_styled_red_bold(self):
        """【待填写：...】 placeholders should be rendered in red bold font."""
        sections = {
            "第一章 投标函": (
                "致招标人：\n"
                "投标人：【待填写：投标人全称】\n"
                "法定代表人：【待填写：法定代表人姓名】\n"
                "项目编号：【待填写：项目编号】\n"
                "我方郑重承诺按照招标文件要求提供全部服务。"
            ),
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_placeholder.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            from docx.shared import RGBColor
            doc = Document(export_path)

            # Find runs that contain placeholder text
            placeholder_runs = []
            normal_runs = []
            for para in doc.paragraphs:
                for run in para.runs:
                    if "待填写" in run.text:
                        placeholder_runs.append(run)
                    elif run.text.strip():
                        normal_runs.append(run)

            # Should have found 3 placeholder runs
            assert len(placeholder_runs) >= 3, (
                f"Expected ≥3 placeholder runs, got {len(placeholder_runs)}: "
                f"{[r.text for r in placeholder_runs]}"
            )

            # Each placeholder run should be bold and red
            for run in placeholder_runs:
                assert run.bold is True, (
                    f"Placeholder '{run.text}' should be bold"
                )
                assert run.font.color.rgb == RGBColor(0xFF, 0x00, 0x00), (
                    f"Placeholder '{run.text}' should be red, got {run.font.color.rgb}"
                )

            # Normal runs should NOT be bold or red
            for run in normal_runs:
                if run.bold is True and run.font.color.rgb == RGBColor(0xFF, 0x00, 0x00):
                    pytest.fail(
                        f"Normal text '{run.text}' should not be red bold"
                    )

    def test_placeholder_mixed_with_normal_text(self):
        """Placeholders mixed inline with normal text should split correctly."""
        sections = {
            "测试章": "本公司【待填写：公司名称】郑重承诺，项目编号【待填写：项目编号】已确认。",
        }
        with tempfile.TemporaryDirectory() as tmpdir:
            export_path = str(Path(tmpdir) / "test_mixed.docx")
            _build_docx(sections, DEFAULT_FORMAT, export_path)

            from docx import Document
            from docx.shared import RGBColor
            doc = Document(export_path)

            # Find the paragraph with mixed content
            mixed_para = None
            for para in doc.paragraphs:
                if "待填写" in para.text and "郑重承诺" in para.text:
                    mixed_para = para
                    break

            assert mixed_para is not None, "Should find paragraph with mixed content"

            # Should have multiple runs: normal, placeholder, normal, placeholder, normal
            runs = mixed_para.runs
            assert len(runs) >= 5, (
                f"Expected ≥5 runs (normal+placeholder+normal+placeholder+normal), "
                f"got {len(runs)}: {[r.text for r in runs]}"
            )

            # Verify placeholder runs are red bold
            placeholder_runs = [r for r in runs if "待填写" in r.text]
            assert len(placeholder_runs) == 2
            for run in placeholder_runs:
                assert run.bold is True
                assert run.font.color.rgb == RGBColor(0xFF, 0x00, 0x00)

            # Verify normal runs are not red bold
            normal_runs = [r for r in runs if "待填写" not in r.text and r.text.strip()]
            for run in normal_runs:
                assert not (run.bold is True and run.font.color.rgb == RGBColor(0xFF, 0x00, 0x00)), \
                    f"Normal text should not be red bold: '{run.text}'"


# ── Node Tests ─────────────────────────────────────────────────────────


class TestDocAssemblerNode:
    def test_empty_sections(self):
        state = factory_state()
        result = doc_assembler(state)
        assert result["export_path"] == ""
        assert result["node_status"]["DocumentAssembler"] == NodeStatus.FAILED.value

    def test_node_returns_export_path(self, state_with_sections):
        with tempfile.TemporaryDirectory() as tmpdir:
            result = doc_assembler(state_with_sections, export_dir=tmpdir)
            assert result["export_path"] != ""
            assert Path(result["export_path"]).exists()
            assert result["node_status"]["DocumentAssembler"] == NodeStatus.COMPLETED.value

    def test_filename_includes_round(self, state_with_sections):
        with tempfile.TemporaryDirectory() as tmpdir:
            result = doc_assembler(state_with_sections, export_dir=tmpdir)
            filename = Path(result["export_path"]).name
            assert "v0" in filename  # current_round=0
            assert filename.endswith(".docx")

    def test_export_dir_auto_created(self, state_with_sections):
        with tempfile.TemporaryDirectory() as tmpdir:
            auto_dir = f"{tmpdir}/auto_created_subdir"
            result = doc_assembler(state_with_sections, export_dir=auto_dir)
            assert Path(result["export_path"]).exists()
